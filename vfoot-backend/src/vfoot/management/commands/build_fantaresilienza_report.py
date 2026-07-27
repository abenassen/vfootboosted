"""Replay a real fantacalcio league season with OUR voto puro and see how the
standings would have changed — matchday by matchday.

    python manage.py build_fantaresilienza_report
    python manage.py build_fantaresilienza_report --dir <folder> --out <file.html>

The league (a Mantra head-to-head league the user actually played) exported one
xlsx per matchday: every team's lineup with the fantacalcio *voto* (base) and
*fantavoto* (base + bonus/malus). The league's own scoring is reproduced exactly
from the sheet — validated: Σ(fielded fantavoti) + defense-modifier + home-factor
== the printed TOTALE on all no-substitution lineups, and TOTALE -> goals (66=1,
+6/goal) reproduces every printed scoreline.

To recompute a season with our votes WITHOUT reinventing the two league rules we
cannot reproduce from scratch (Mantra's module auto-completion on substitutions,
and the defense-modifier bracket table), we anchor to the sheet and swap ONLY the
base vote:

    new_TOTALE = sheet_TOTALE + Σ_fielded (our_voto_puro - fanta_voto)

The bonus/malus (goals, assists, cards) are facts, identical in both worlds, so
they cancel; the defense modifier and home factor are inherited unchanged. The
fielded eleven are the played starters plus the bench substitutes the app used —
recovered by the unique earliest-bench subset whose fantavoti reproduce the
printed TOTALE (validated on 242/244 substitution lineups).

Player -> our DB id reuses the fantacalcio name matcher (surname + first-name
initial, which fantacalcio keeps unique). League matchday L maps to Serie A
matchday L (the league mirrors the first 36 rounds).

The report is a self-contained HTML page: each team's league-position trajectory
(real vs our-votes) over the 36 matchdays, plus the two final tables side by side.
Re-run after any weight change to regenerate it.
"""
from __future__ import annotations

import glob
import html
import re
from collections import defaultdict
from itertools import combinations
from pathlib import Path

import openpyxl
from django.conf import settings
from django.core.management.base import BaseCommand, CommandError

from realdata.models import Match, MatchAppearance
from vfoot.management.commands.voto_puro_discrepancies import Command as DiscCmd
from vfoot.services.classic_pagella import (
    get_reference, get_role_averages, pagella_for_match,
)

SCORE_RE = re.compile(r"^\s*\d+\s*-\s*\d+\s*$")
GD_RE = re.compile(r"_(\d+)_giornata")
DEFAULT_DIR = str(Path(settings.BASE_DIR).parent.parent / "fantaresilienza2025-2026")


def _num(v):
    if v in (None, "", "-"):
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def parse_file(path):
    """One matchday sheet -> list of matchups. Each side: starters/bench lists of
    {role, name, voto, fanta} and a trailer {moddif, fattore, totale}."""
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    def cell(r, c):
        return rows[r][c] if r < len(rows) and c < len(rows[r]) else None

    matchups = []
    i, n = 0, len(rows)
    while i < n:
        c5 = cell(i, 5)
        if c5 is not None and SCORE_RE.match(str(c5)):
            mu = {"teamA": cell(i, 0), "teamB": cell(i, 6), "score": str(c5).strip(),
                  "starters": {"A": [], "B": []}, "bench": {"A": [], "B": []},
                  "trailer": {"A": {}, "B": {}}}
            i += 1  # module row (unused)
            i += 1
            phase = "starters"
            while i < n:
                if (nc := cell(i, 5)) is not None and SCORE_RE.match(str(nc)):
                    break  # next matchup
                s0 = str(cell(i, 0) or "").strip()
                s6 = str(cell(i, 6) or "").strip()
                if s0.startswith("Panchina"):
                    phase = "bench"
                    i += 1
                    continue
                trailer_row = s0.startswith(("Modificatore", "Fattore", "TOTALE", "Inserita")) or \
                    s6.startswith(("Modificatore", "Fattore", "TOTALE", "Inserita"))
                if trailer_row:
                    for side, lc, vc in (("A", 0, 4), ("B", 6, 10)):
                        lab = str(cell(i, lc) or "").strip()
                        val = cell(i, vc)
                        if lab.startswith("Modificatore"):
                            mu["trailer"][side]["moddif"] = _num(val)
                        elif lab.startswith("Fattore"):
                            mu["trailer"][side]["fattore"] = _num(val)
                        elif lab.startswith("TOTALE"):
                            m = re.search(r"([\d.,]+)", lab)
                            if m:
                                mu["trailer"][side]["totale"] = float(m.group(1).replace(",", "."))
                    i += 1
                    continue
                for side, rc, nc2, vc, fc in (("A", 0, 1, 3, 4), ("B", 6, 7, 9, 10)):
                    role = str(cell(i, rc) or "").strip()
                    name = cell(i, nc2)
                    if name and role not in ("", "Panchina"):
                        rec = {"role": role, "name": str(name).strip(),
                               "voto": _num(cell(i, vc)), "fanta": _num(cell(i, fc))}
                        (mu["starters"] if phase == "starters" else mu["bench"])[side].append(rec)
                i += 1
            matchups.append(mu)
        else:
            i += 1
    return matchups


def goals_from_total(t):
    """League fantapunti -> goals: 66=1, then +6 per goal."""
    return 0 if t < 66 else int((t - 66) // 6) + 1


class Command(BaseCommand):
    help = "Replay a fantacalcio league season with our voto puro; compare standings."

    def add_arguments(self, parser):
        parser.add_argument("--season", type=int, default=2)
        parser.add_argument("--dir", default=DEFAULT_DIR,
                            help="Folder of per-matchday league xlsx files.")
        parser.add_argument("--out", default=str(Path(settings.BASE_DIR).parent.parent
                                                 / "fantaresilienza_report.html"))

    # -- player id resolution (club-less: surname + fantacalcio initial) --------
    def _build_name_index(self, cs_id):
        disc = DiscCmd()
        surn2pid = defaultdict(set)
        pid_first = defaultdict(set)
        for pid, full, short in (MatchAppearance.objects
                                 .filter(match__competition_season_id=cs_id)
                                 .values_list("player_id", "player__full_name",
                                              "player__short_name").distinct()):
            for nm in (full, short):
                s, _ = disc._surname_and_initial(nm)
                if s:
                    surn2pid[s].add(pid)
            ft = disc._afold(full or "").split()
            if ft:
                pid_first[pid].add(ft[0])

        def resolve(name):
            s, ab = disc._surname_and_initial(name)
            c = surn2pid.get(s, set())
            if len(c) == 1:
                return next(iter(c))
            if len(c) > 1 and ab:
                nar = [pid for pid in c if any(fn.startswith(ab) for fn in pid_first.get(pid, ()))]
                if len(nar) == 1:
                    return nar[0]
            return None
        return resolve

    def _our_voto_puro(self, cs_id):
        ref = get_reference(cs_id)
        avgs = get_role_averages(cs_id)
        vp = {}
        for m in Match.objects.filter(competition_season_id=cs_id):
            pag = pagella_for_match(m, ref, averages=avgs)
            for side in ("home", "away"):
                for ln in pag[side]["starters"] + pag[side]["bench"]:
                    if ln.get("voto_puro") is not None:
                        vp[(m.matchday, ln["player_id"])] = float(ln["voto_puro"])
        return vp

    @staticmethod
    def _find_subs(bench, gap, k):
        cand = [(bi, b) for bi, b in enumerate(bench) if b["fanta"] is not None]
        best = None
        for combo in combinations(range(len(cand)), k):
            if abs(sum(cand[j][1]["fanta"] for j in combo) - gap) < 0.01:
                key = tuple(cand[j][0] for j in combo)
                if best is None or key < best[0]:
                    best = (key, [cand[j][1] for j in combo])
        return best[1] if best else None

    def handle(self, *args, **opt):
        cs_id = opt["season"]
        files = sorted(glob.glob(f"{opt['dir']}/*.xlsx"),
                       key=lambda p: int(GD_RE.search(p).group(1)) if GD_RE.search(p) else 0)
        if not files:
            raise CommandError(f"No league xlsx in {opt['dir']}")

        self.stdout.write("Computing our voto puro for every match…")
        our_vp = self._our_voto_puro(cs_id)
        resolve = self._build_name_index(cs_id)

        stats = dict(voted=0, unmatched=0, sub_exact=0, sub_approx=0, goal_ok=0, goal_bad=0)
        # records[gd] = list of (teamA, teamB, gAa, gBa, gAo, gBo)
        records = defaultdict(list)
        for f in files:
            L = int(GD_RE.search(f).group(1))
            for mu in parse_file(f):
                side_goals = {}
                for side in ("A", "B"):
                    st, bench, tr = mu["starters"][side], mu["bench"][side], mu["trailer"][side]
                    tot = tr.get("totale")
                    if tot is None:
                        continue
                    mod = tr.get("moddif") or 0
                    fat = tr.get("fattore") or 0
                    played = [p for p in st if p["fanta"] is not None]
                    svs = [p for p in st if p["fanta"] is None]
                    fielded = list(played)
                    if svs:
                        gap = round(tot - sum(p["fanta"] for p in played) - mod - fat, 1)
                        subs = self._find_subs(bench, gap, len(svs))
                        if subs is not None:
                            stats["sub_exact"] += 1
                        else:
                            stats["sub_approx"] += 1
                            subs = [b for b in bench if b["fanta"] is not None][:len(svs)]
                        fielded += subs
                    delta = 0.0
                    for p in fielded:
                        stats["voted"] += 1
                        pid = resolve(p["name"])
                        vp = our_vp.get((L, pid)) if pid else None
                        if vp is None:
                            stats["unmatched"] += 1
                            continue
                        delta += vp - p["voto"]
                    side_goals[side] = {"act": goals_from_total(tot),
                                        "our": goals_from_total(tot + delta)}
                if len(side_goals) != 2:
                    continue
                gAs, gBs = (int(x) for x in re.split(r"\s*-\s*", mu["score"]))
                if side_goals["A"]["act"] == gAs and side_goals["B"]["act"] == gBs:
                    stats["goal_ok"] += 1
                else:
                    stats["goal_bad"] += 1
                records[L].append((mu["teamA"], mu["teamB"],
                                   side_goals["A"]["act"], side_goals["B"]["act"],
                                   side_goals["A"]["our"], side_goals["B"]["our"]))

        gds = sorted(records)
        teams = sorted({t for L in gds for r in records[L] for t in (r[0], r[1])})
        pos_act, pts_act = self._trajectory(records, gds, teams, actual=True)
        pos_our, pts_our = self._trajectory(records, gds, teams, actual=False)

        html_str = self._render(gds, teams, records, pos_act, pos_our,
                                 pts_act, pts_our, stats)
        Path(opt["out"]).write_text(html_str, encoding="utf-8")

        self.stdout.write(self.style.SUCCESS(
            f"\nfielded&voted players {stats['voted']}  "
            f"unmatched {stats['unmatched']} ({100*stats['unmatched']/stats['voted']:.1f}%)  "
            f"subs exact {stats['sub_exact']}/{stats['sub_exact']+stats['sub_approx']}  "
            f"goal-conversion {stats['goal_ok']}/{stats['goal_ok']+stats['goal_bad']}"))
        self.stdout.write(f"report -> {opt['out']}")

    # -- standings -------------------------------------------------------------
    @staticmethod
    def _standings_after(records, gds_upto, actual):
        pts = defaultdict(int); gf = defaultdict(int); ga = defaultdict(int)
        for L in gds_upto:
            for tA, tB, gAa, gBa, gAo, gBo in records[L]:
                gA, gB = (gAa, gBa) if actual else (gAo, gBo)
                gf[tA] += gA; ga[tA] += gB; gf[tB] += gB; ga[tB] += gA
                if gA > gB:
                    pts[tA] += 3
                elif gB > gA:
                    pts[tB] += 3
                else:
                    pts[tA] += 1; pts[tB] += 1
        order = sorted(pts, key=lambda t: (-pts[t], -(gf[t] - ga[t]), -gf[t]))
        return order, pts, gf, ga

    def _trajectory(self, records, gds, teams, actual):
        pos = {t: [] for t in teams}
        final_pts = {}
        for k in range(1, len(gds) + 1):
            order, pts, gf, ga = self._standings_after(records, gds[:k], actual)
            rank = {t: i + 1 for i, t in enumerate(order)}
            for t in teams:
                pos[t].append(rank.get(t))
            final_pts = pts
        return pos, final_pts

    # -- HTML ------------------------------------------------------------------
    def _render(self, gds, teams, records, pos_act, pos_our, pts_act, pts_our, stats):
        N = len(gds)
        act_order, a_pts, a_gf, a_ga = self._standings_after(records, gds, True)
        our_order, o_pts, o_gf, o_ga = self._standings_after(records, gds, False)
        a_rank = {t: i + 1 for i, t in enumerate(act_order)}
        o_rank = {t: i + 1 for i, t in enumerate(our_order)}

        def esc(s):
            return html.escape(str(s))

        # small-multiple SVG per team: position (1 top .. n bottom) over matchdays
        def spark(team):
            w, h, pad = 260, 120, 18
            n_teams = len(teams)
            def X(i):
                return pad + (w - 2 * pad) * i / max(N - 1, 1)
            def Y(p):
                return pad + (h - 2 * pad) * (p - 1) / max(n_teams - 1, 1)
            def line(series, cls):
                pts = " ".join(f"{X(i):.1f},{Y(p):.1f}" for i, p in enumerate(series) if p)
                return f'<polyline class="{cls}" points="{pts}"/>'
            fa, fo = a_rank[team], o_rank[team]
            return (
                f'<svg viewBox="0 0 {w} {h}" class="spark" role="img">'
                f'<line class="axis" x1="{pad}" y1="{Y(1):.1f}" x2="{w-pad}" y2="{Y(1):.1f}"/>'
                f'<line class="axis" x1="{pad}" y1="{Y(n_teams):.1f}" x2="{w-pad}" y2="{Y(n_teams):.1f}"/>'
                f'{line(pos_act[team], "l-act")}{line(pos_our[team], "l-our")}'
                f'<circle class="d-act" cx="{X(N-1):.1f}" cy="{Y(fa):.1f}" r="3"/>'
                f'<circle class="d-our" cx="{X(N-1):.1f}" cy="{Y(fo):.1f}" r="3"/>'
                f'</svg>')

        cards = []
        for t in sorted(teams, key=lambda x: o_rank[x]):
            d = a_rank[t] - o_rank[t]
            badge = ("=" if d == 0 else (f"▲{d}" if d > 0 else f"▼{-d}"))
            bcls = "flat" if d == 0 else ("up" if d > 0 else "down")
            cards.append(
                f'<figure class="card"><figcaption><span class="tname">{esc(t)}</span>'
                f'<span class="delta {bcls}">{badge}</span></figcaption>{spark(t)}'
                f'<div class="legend"><span><i class="s-act"></i>reale {a_rank[t]}º·{a_pts[t]}pt</span>'
                f'<span><i class="s-our"></i>nostri {o_rank[t]}º·{o_pts[t]}pt</span></div></figure>')

        # side-by-side final tables
        def table(order, pts, gf, ga, rank_other, title, klass):
            rows = []
            for i, t in enumerate(order, 1):
                d = rank_other[t] - i
                mv = "" if d == 0 else (f'<span class="up">▲{d}</span>' if d > 0
                                        else f'<span class="down">▼{-d}</span>')
                rows.append(f'<tr><td class="pos">{i}</td><td>{esc(t)}</td>'
                            f'<td class="pt">{pts[t]}</td><td class="gd">{gf[t]-ga[t]:+d}</td>'
                            f'<td class="mv">{mv}</td></tr>')
            return (f'<table class="{klass}"><caption>{title}</caption><thead><tr>'
                    f'<th>#</th><th>squadra</th><th>pt</th><th>±</th><th>Δ</th></tr></thead>'
                    f'<tbody>{"".join(rows)}</tbody></table>')

        t_act = table(act_order, a_pts, a_gf, a_ga, o_rank, "Classifica reale (voti fanta)", "real")
        t_our = table(our_order, o_pts, o_gf, o_ga, a_rank, "Classifica con i nostri voti", "our")

        cov = 100 * stats["unmatched"] / max(stats["voted"], 1)
        subtot = stats["sub_exact"] + stats["sub_approx"]
        note = (f"{stats['voted']} giocatori schierati · copertura voto puro "
                f"{100-cov:.1f}% · sostituzioni ricostruite {stats['sub_exact']}/{subtot} "
                f"· conversione punti→gol {stats['goal_ok']}/{stats['goal_ok']+stats['goal_bad']} ✓")

        # headline: the title change + the biggest position swings
        champ_a, champ_o = act_order[0], our_order[0]
        movers = sorted(teams, key=lambda t: -abs(a_rank[t] - o_rank[t]))
        chips = []
        if champ_a != champ_o:
            chips.append(f'<span class="chip win"><b>{esc(champ_o)}</b> scippa il titolo a '
                         f'{esc(champ_a)}</span>')
        else:
            chips.append(f'<span class="chip win"><b>{esc(champ_o)}</b> resta campione</span>')
        for t in movers[:3]:
            d = a_rank[t] - o_rank[t]
            if d == 0:
                continue
            cls = "up" if d > 0 else "down"
            arr = f"▲{d}" if d > 0 else f"▼{-d}"
            chips.append(f'<span class="chip {cls}"><b>{esc(t)}</b> {a_rank[t]}º→{o_rank[t]}º '
                         f'<em>{arr}</em></span>')
        lead = "".join(chips)

        return _PAGE.replace("__CARDS__", "\n".join(cards)) \
                    .replace("__TREAL__", t_act).replace("__TOUR__", t_our) \
                    .replace("__LEAD__", lead) \
                    .replace("__NOTE__", esc(note))


_PAGE = """<title>Fantaresilienza · voti fanta vs nostri voti</title>
<style>
/* warm-neutral ground, hue-biased toward the pitch-orange accent */
:root{--bg:#fbfaf8;--fg:#20211d;--mut:#78756c;--line:#e7e4dd;--card:#fff;
--act:#2b2a26;--our:#d1541f;--up:#2f7d47;--down:#bd3d2c;--flat:#a8a49a;}
@media(prefers-color-scheme:dark){:root{--bg:#12120f;--fg:#eceae4;--mut:#98948a;
--line:#2a2823;--card:#1a1915;--act:#cfccc3;--our:#ff7a45;--up:#59c07d;--down:#e8695a;}}
:root[data-theme=dark]{--bg:#12120f;--fg:#eceae4;--mut:#98948a;--line:#2a2823;
--card:#1a1915;--act:#cfccc3;--our:#ff7a45;--up:#59c07d;--down:#e8695a;}
:root[data-theme=light]{--bg:#fbfaf8;--fg:#20211d;--mut:#78756c;--line:#e7e4dd;
--card:#fff;--act:#2b2a26;--our:#d1541f;--up:#2f7d47;--down:#bd3d2c;--flat:#a8a49a;}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--fg);padding:30px 30px 48px;
font:15px/1.5 ui-sans-serif,-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif;
font-variant-numeric:tabular-nums;-webkit-font-smoothing:antialiased}
.wrap{max-width:1120px;margin:0 auto}
h1{font-size:23px;line-height:1.2;margin:0 0 4px;letter-spacing:-.2px;text-wrap:balance}
.sub{color:var(--mut);margin:0 0 18px;font-size:13px;max-width:70ch}
.lead{display:flex;flex-wrap:wrap;gap:8px;margin:0 0 26px}
.chip{font-size:12.5px;padding:5px 11px;border-radius:8px;border:1px solid var(--line);background:var(--card)}
.chip b{font-weight:680}.chip em{font-style:normal;font-weight:750;margin-left:2px}
.chip.win{border-color:color-mix(in srgb,var(--our) 45%,var(--line));
background:color-mix(in srgb,var(--our) 9%,var(--card))}
.chip.up em{color:var(--up)}.chip.down em{color:var(--down)}
.eyebrow{text-transform:uppercase;letter-spacing:.9px;font-size:11px;font-weight:670;
color:var(--mut);margin:0 0 12px}
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(255px,1fr));gap:13px;margin-bottom:34px}
.card{margin:0;background:var(--card);border:1px solid var(--line);border-radius:11px;padding:11px 13px 9px}
figcaption{display:flex;justify-content:space-between;align-items:center;gap:6px;margin-bottom:3px}
.tname{font-weight:660;font-size:12px;letter-spacing:.1px;line-height:1.15}
.delta{font-weight:730;font-size:11.5px;padding:1px 7px;border-radius:20px;white-space:nowrap}
.delta.up{color:var(--up);background:color-mix(in srgb,var(--up) 14%,transparent)}
.delta.down{color:var(--down);background:color-mix(in srgb,var(--down) 14%,transparent)}
.delta.flat{color:var(--flat)}
.spark{width:100%;height:auto;display:block}
.spark .axis{stroke:var(--line);stroke-width:1}
.spark polyline{fill:none;stroke-width:2;stroke-linejoin:round;stroke-linecap:round}
.spark .l-act{stroke:var(--act);opacity:.82}
.spark .l-our{stroke:var(--our)}
.spark .d-act{fill:var(--act)}.spark .d-our{fill:var(--our)}
.legend{display:flex;gap:12px;font-size:10.5px;color:var(--mut);margin-top:1px}
.legend i{display:inline-block;width:11px;height:0;border-top:2px solid;vertical-align:middle;margin-right:4px}
.legend .s-act{border-color:var(--act)}.legend .s-our{border-color:var(--our)}
.tables{display:flex;gap:24px;flex-wrap:wrap}
table{border-collapse:collapse;font-size:13.5px;flex:1;min-width:308px}
caption{text-align:left;font-weight:680;padding:0 0 9px;font-size:13.5px;letter-spacing:-.1px}
th{text-align:left;color:var(--mut);font-weight:620;font-size:11px;text-transform:uppercase;
letter-spacing:.4px;padding:4px 9px;border-bottom:1.5px solid var(--line)}
td{padding:6px 9px;border-bottom:1px solid var(--line)}
tr:last-child td{border-bottom:none}
td.pos{color:var(--mut);width:26px}td.pt{font-weight:730;text-align:right}
td.gd{text-align:right;color:var(--mut)}
td.mv{text-align:right;width:42px}
.our caption{color:var(--our)}
.up{color:var(--up);font-weight:730}.down{color:var(--down);font-weight:730}
.foot{color:var(--mut);font-size:11.5px;margin-top:26px;max-width:80ch}
</style>
<div class="wrap">
<h1>Fantaresilienza 2025-26 — come sarebbe andata con i nostri voti</h1>
<p class="sub">Stessa lega, stesse formazioni, stessi bonus/malus: sostituito solo il
<b>voto base</b> di ogni giocatore con il nostro voto puro, e ricalcolati punteggi,
risultati e classifica di ogni giornata.</p>
<div class="lead">__LEAD__</div>
<p class="eyebrow">Posizione giornata per giornata · linea scura = reale · <span style="color:var(--our)">arancio = nostri voti</span> · 1º in alto</p>
<div class="grid">__CARDS__</div>
<div class="tables">__TREAL__ __TOUR__</div>
<p class="foot">__NOTE__</p>
</div>
"""
